import os
import openpyxl

from core.banco import conecta


# ============================================================
# CONFIGURAÇÃO
# ============================================================

arquivo = os.path.join(
    os.path.expanduser("~"),
    "Desktop",
    "TROCA.xlsx"
)


# ============================================================
# PROCESSAMENTO
# ============================================================

def processar_excel():
    try:
        print(f"Lendo arquivo: {arquivo}")

        # Abre o Excel existente
        wb = openpyxl.load_workbook(arquivo)

        # Aba original
        ws_origem = wb["Planilha1"]

        # Se já existir, remove a aba Produtos
        if "Produtos" in wb.sheetnames:
            del wb["Produtos"]

        # Cria nova aba
        ws_destino = wb.create_sheet("Produtos")

        # Cabeçalho
        cabecalho = [
            "Cod Filho",
            "Desc Filho",
            "Ref Filho",
            "UM Filho",
            "Qtde",
            "Conjunto",
            "Custo Unitário",
            "Tipo"
        ]

        ws_destino.append(cabecalho)

        # Cursor do banco
        cursor = conecta.cursor()

        # Percorre os produtos da planilha
        for linha in ws_origem.iter_rows(min_row=2, values_only=True):

            cod_produto = linha[0]
            descricao = linha[1]
            referencia = linha[2]
            unidade = linha[3]
            quantidade = linha[4]

            if not cod_produto:
                continue

            print(f"Consultando produto {cod_produto}...")

            # Consulta banco
            cursor.execute(
                f"""
                SELECT prod.codigo, conj.conjunto, prod.custounitario, tip.TIPOMATERIAL 
                FROM produto as prod
                LEFT JOIN tipomaterial as tip ON prod.tipomaterial = tip.id 
                INNER JOIN conjuntos conj ON prod.conjunto = conj.id 
                WHERE prod.codigo = {cod_produto}
                """
            )

            resultado = cursor.fetchall()

            # Produto encontrado
            if resultado:
                cod, conjunto, custo, tipo = resultado[0]
            else:
                conjunto = None
                custo = None
                tipo = None

            # Adiciona na nova aba
            ws_destino.append([
                cod_produto,
                descricao,
                referencia,
                unidade,
                quantidade,
                conjunto,
                custo,
                tipo
            ])

        cursor.close()

        # Ajusta largura das colunas
        larguras = {
            "A": 15,
            "B": 45,
            "C": 20,
            "D": 12,
            "E": 15,
            "F": 20,
            "G": 18
        }

        for coluna, largura in larguras.items():
            ws_destino.column_dimensions[coluna].width = largura

        # Congela cabeçalho
        ws_destino.freeze_panes = "A2"

        # Salva no MESMO arquivo
        wb.save(arquivo)

        print()
        print("==========================================")
        print("Processamento concluído!")
        print(f"Arquivo atualizado: {arquivo}")
        print("Nova aba criada: Produtos")
        print("==========================================")

    except Exception as e:
        print()
        print("ERRO:")
        print(e)


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    processar_excel()